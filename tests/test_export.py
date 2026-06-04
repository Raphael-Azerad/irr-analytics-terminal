"""Tests for Excel export."""

from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from irr_terminal.calculations import calculate_metrics
from irr_terminal.comparison import compare_projects
from irr_terminal.export import SHEET_NAMES, build_excel_workbook
from irr_terminal.scenarios import scenario_analysis
from irr_terminal.sensitivity import npv_profile, sensitivity_table


def test_excel_export_workbook_structure() -> None:
    flows = [-1000, 400, 500, 600]
    cash_frame = pd.DataFrame({"Period": range(4), "Cash Flow": flows})
    workbook = build_excel_workbook(
        "Test",
        cash_frame,
        calculate_metrics(flows, 0.10, 0.10),
        npv_profile(flows),
        scenario_analysis(flows, 0.10, 0.10),
        sensitivity_table(flows, [0.10], [0.0]),
        compare_projects({"Test": flows}, 0.10, 0.10),
        0.10,
        0.10,
    )
    assert load_workbook(BytesIO(workbook)).sheetnames == SHEET_NAMES
