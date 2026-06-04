"""Excel and PDF investment committee report generation."""

from __future__ import annotations

from io import BytesIO
from datetime import date, datetime
from typing import Iterable

import pandas as pd
from fpdf import FPDF
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from calculations.metrics import FinancialMetrics, xirr, xnpv
from calculations.recommendation import InvestmentRecommendation


def build_excel_report(
    project_name: str,
    investment_name: str,
    cash_flows: Iterable[float],
    metrics: FinancialMetrics,
    recommendation: InvestmentRecommendation,
    scenarios: pd.DataFrame,
    dates: Iterable[date | datetime | str] | None = None,
) -> bytes:
    """Build a multi-tab professional Excel workbook."""

    output = BytesIO()
    flows = list(cash_flows)
    date_values = list(dates) if dates is not None else None
    metric_names = [
        "Initial Investment",
        "Total Cash Inflows",
        "IRR",
        "NPV",
        "MIRR",
        "Profitability Index",
        "Payback Period",
        "Discounted Payback Period",
        "IRR Spread to Hurdle",
        "NPV / Initial Investment",
    ]
    metric_values: list[object] = [
        metrics.initial_investment,
        metrics.total_inflows,
        metrics.irr,
        metrics.npv,
        metrics.mirr,
        metrics.profitability_index,
        metrics.payback_period,
        metrics.discounted_payback_period,
        metrics.irr_spread_to_hurdle,
        metrics.npv_to_investment,
    ]
    if date_values is not None:
        metric_names.extend(["XIRR", "XNPV"])
        metric_values.extend(
            [
                xirr(flows, date_values),
                xnpv(metrics.discount_rate, flows, date_values),
            ]
        )
    metric_names.append("Recommendation")
    metric_values.append(recommendation.rating)
    metrics_frame = pd.DataFrame({"Metric": metric_names, "Value": metric_values})

    discounted_flows = [
        flow / (1 + metrics.discount_rate) ** period
        for period, flow in enumerate(flows)
    ]
    cash_flow_data: dict[str, object] = {
        "Period": range(len(flows)),
        "Cash Flow": flows,
        "Cumulative Cash Flow": pd.Series(flows).cumsum(),
        "Discounted Cash Flow": discounted_flows,
        "Cumulative Discounted Cash Flow": pd.Series(discounted_flows).cumsum(),
    }
    if date_values is not None:
        cash_flow_data["Date"] = date_values
    cash_flow_frame = pd.DataFrame(cash_flow_data)
    assumptions = pd.DataFrame(
        {
            "Assumption": ["Project Name", "Investment Name", "Discount Rate", "Hurdle Rate", "Cost of Capital"],
            "Value": [project_name, investment_name, metrics.discount_rate, metrics.hurdle_rate, metrics.cost_of_capital],
        }
    )
    decision_matrix = pd.DataFrame(
        {
            "Test": [
                "NPV creates value",
                "IRR clears hurdle rate",
                "Profitability Index exceeds 1.0x",
                "Bear-case NPV creates value",
            ],
            "Result": [
                metrics.npv,
                metrics.irr,
                metrics.profitability_index,
                float(scenarios.loc[scenarios["Scenario"] == "Bear", "NPV"].iloc[0]),
            ],
            "Status": [
                "Pass" if metrics.npv > 0 else "Fail",
                "Pass" if metrics.irr is not None and metrics.irr > metrics.hurdle_rate else "Fail",
                "Pass" if metrics.profitability_index is not None and metrics.profitability_index > 1 else "Fail",
                "Pass" if float(scenarios.loc[scenarios["Scenario"] == "Bear", "NPV"].iloc[0]) > 0 else "Diligence",
            ],
        }
    )
    rationale = pd.DataFrame({"Investment Committee Rationale": recommendation.rationale})

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        metrics_frame.to_excel(writer, sheet_name="Executive Summary", index=False)
        cash_flow_frame.to_excel(writer, sheet_name="Cash Flows", index=False)
        scenarios.to_excel(writer, sheet_name="Scenario Analysis", index=False)
        decision_matrix.to_excel(writer, sheet_name="Decision Matrix", index=False)
        rationale.to_excel(writer, sheet_name="Recommendation", index=False)
        assumptions.to_excel(writer, sheet_name="Assumptions", index=False)

        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.sheet_view.showGridLines = False
            for cell in worksheet[1]:
                cell.fill = PatternFill("solid", fgColor="0B1B2B")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center")
            for column in worksheet.columns:
                width = max(len(str(cell.value or "")) for cell in column) + 3
                worksheet.column_dimensions[get_column_letter(column[0].column)].width = min(width, 36)
    return output.getvalue()


def build_pdf_report(
    project_name: str,
    investment_name: str,
    metrics: FinancialMetrics,
    recommendation: InvestmentRecommendation,
    scenarios: pd.DataFrame,
) -> bytes:
    """Build a concise investment memo PDF."""

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_fill_color(8, 17, 31)
    pdf.rect(0, 0, 210, 38, style="F")
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 20)
    pdf.set_xy(15, 12)
    pdf.cell(0, 8, "IRR Analytics Terminal")
    pdf.set_font("Helvetica", "", 11)
    pdf.set_xy(15, 23)
    pdf.cell(0, 6, f"Investment Committee Memo | {project_name}")

    pdf.set_text_color(20, 30, 45)
    pdf.set_y(48)
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 8, investment_name, ln=True)
    pdf.set_font("Helvetica", "B", 13)
    pdf.set_text_color(0, 130, 100)
    pdf.cell(0, 8, f"Recommendation: {recommendation.rating}", ln=True)
    pdf.set_text_color(20, 30, 45)
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(0, 6, recommendation.headline)

    pdf.ln(4)
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "Executive Summary", ln=True)
    rows = [
        ("Initial Investment", f"${metrics.initial_investment:,.0f}"),
        ("IRR", "N/A" if metrics.irr is None else f"{metrics.irr:.2%}"),
        ("NPV", f"${metrics.npv:,.0f}"),
        ("MIRR", "N/A" if metrics.mirr is None else f"{metrics.mirr:.2%}"),
        ("Profitability Index", "N/A" if metrics.profitability_index is None else f"{metrics.profitability_index:.2f}x"),
        ("Payback Period", "N/A" if metrics.payback_period is None else f"{metrics.payback_period:.2f} years"),
    ]
    for label, value in rows:
        pdf.set_font("Helvetica", "B", 10)
        pdf.cell(65, 7, label)
        pdf.set_font("Helvetica", "", 10)
        pdf.cell(0, 7, value, ln=True)

    pdf.ln(4)
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "Investment Rationale", ln=True)
    pdf.set_font("Helvetica", "", 10)
    for reason in recommendation.rationale:
        pdf.multi_cell(0, 6, f"- {reason}")

    pdf.ln(4)
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "Scenario Analysis", ln=True)
    pdf.set_font("Helvetica", "B", 9)
    pdf.cell(35, 7, "Scenario")
    pdf.cell(35, 7, "IRR")
    pdf.cell(50, 7, "NPV", ln=True)
    pdf.set_font("Helvetica", "", 9)
    for _, row in scenarios.iterrows():
        pdf.cell(35, 7, str(row["Scenario"]))
        irr_value = row["IRR"]
        pdf.cell(35, 7, "N/A" if pd.isna(irr_value) else f"{irr_value:.2%}")
        pdf.cell(50, 7, f"${row['NPV']:,.0f}", ln=True)

    pdf.ln(4)
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "Return Benchmark Chart", ln=True)
    chart_x, chart_y, chart_width = 20, pdf.get_y() + 4, 125
    max_rate = max(
        metrics.irr or 0,
        metrics.hurdle_rate,
        metrics.cost_of_capital,
        0.01,
    )
    benchmarks = [
        ("IRR", metrics.irr or 0, (40, 180, 190)),
        ("Hurdle Rate", metrics.hurdle_rate, (244, 201, 93)),
        ("Cost of Capital", metrics.cost_of_capital, (61, 220, 151)),
    ]
    for index, (label, value, color) in enumerate(benchmarks):
        y = chart_y + index * 12
        pdf.set_font("Helvetica", "", 9)
        pdf.set_text_color(20, 30, 45)
        pdf.text(chart_x, y + 5, label)
        pdf.set_fill_color(*color)
        pdf.rect(chart_x + 35, y, chart_width * value / max_rate, 6, style="F")
        pdf.text(chart_x + 38 + chart_width * value / max_rate, y + 5, f"{value:.2%}")
    pdf.set_y(chart_y + 42)

    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "Key Assumptions", ln=True)
    pdf.set_font("Helvetica", "", 10)
    for label, value in [
        ("Discount Rate", f"{metrics.discount_rate:.2%}"),
        ("Hurdle Rate", f"{metrics.hurdle_rate:.2%}"),
        ("Cost of Capital", f"{metrics.cost_of_capital:.2%}"),
    ]:
        pdf.cell(65, 6, label)
        pdf.cell(0, 6, value, ln=True)

    bear_npv = float(scenarios.loc[scenarios["Scenario"] == "Bear", "NPV"].iloc[0])
    pdf.ln(3)
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 8, "Risk Analysis", ln=True)
    pdf.set_font("Helvetica", "", 10)
    if bear_npv < 0:
        risk_text = (
            "The bear case produces a negative NPV, indicating that value creation "
            "is sensitive to adverse operating performance. Investment approval "
            "should be paired with diligence on revenue resilience, cost controls, "
            "and downside mitigation."
        )
    else:
        risk_text = (
            "The project maintains a positive NPV in the bear case, indicating a "
            "degree of downside resilience. Diligence should still focus on the "
            "operating assumptions that drive the scenario range."
        )
    pdf.multi_cell(0, 6, risk_text)

    return bytes(pdf.output())
