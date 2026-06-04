"""Formatted Excel export for investment analysis."""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .calculations import FinancialMetrics


SHEET_NAMES = [
    "Summary",
    "Cash Flows",
    "Metrics",
    "NPV Profile",
    "Scenario Analysis",
    "Sensitivity Analysis",
    "Multi-Project Comparison",
    "Assumptions",
]


def build_excel_workbook(
    project_name: str,
    cash_flow_frame: pd.DataFrame,
    metrics: FinancialMetrics,
    profile: pd.DataFrame,
    scenarios: pd.DataFrame,
    sensitivity: pd.DataFrame,
    comparison: pd.DataFrame,
    discount_rate: float,
    hurdle_rate: float,
) -> bytes:
    output = BytesIO()
    summary = pd.DataFrame(
        {
            "Project": [project_name],
            "Decision Status": [metrics.decision_status],
            "NPV": [metrics.npv],
            "IRR": [metrics.irr],
            "MIRR": [metrics.mirr],
        }
    )
    metrics_frame = pd.DataFrame(
        {"Metric": list(metrics.to_dict().keys()), "Value": list(metrics.to_dict().values())}
    )
    assumptions = pd.DataFrame(
        {
            "Assumption": ["Discount Rate", "Hurdle Rate"],
            "Value": [discount_rate, hurdle_rate],
        }
    )
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        cash_flow_frame.to_excel(writer, sheet_name="Cash Flows", index=False)
        metrics_frame.to_excel(writer, sheet_name="Metrics", index=False)
        profile.to_excel(writer, sheet_name="NPV Profile", index=False)
        scenarios.to_excel(writer, sheet_name="Scenario Analysis", index=False)
        sensitivity.to_excel(writer, sheet_name="Sensitivity Analysis")
        comparison.to_excel(writer, sheet_name="Multi-Project Comparison", index=False)
        assumptions.to_excel(writer, sheet_name="Assumptions", index=False)
        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.sheet_view.showGridLines = False
            for cell in worksheet[1]:
                cell.fill = PatternFill("solid", fgColor="0B1B2B")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center")
            for column in worksheet.columns:
                width = max(len(str(cell.value or "")) for cell in column) + 3
                worksheet.column_dimensions[get_column_letter(column[0].column)].width = min(
                    width, 36
                )
            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    header = worksheet.cell(row=1, column=cell.column).value
                    if header and any(
                        word in str(header).lower()
                        for word in ["npv", "investment", "cash flow"]
                    ):
                        cell.number_format = '$#,##0'
                    elif header and any(
                        word in str(header).lower()
                        for word in ["irr", "mirr", "rate", "probability"]
                    ):
                        cell.number_format = "0.00%"
    return output.getvalue()
